from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/empirical-study/agentplat-v28-empirical-study-paper-appendix.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5B6470"
LIGHT = "F2F4F7"
CALLOUT = "E8EEF5"
RED = "9B1C1C"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.first_child_found_in("w:tblInd")
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side in ("top", "bottom", "start", "end"):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), "80" if side in ("top", "bottom") else "120")
                node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run(run, size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_paragraph(doc, text="", style=None, *, before=None, after=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def add_label_value(doc, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(label + ": ")
    set_run(r, 10.5, INK, True)
    r = p.add_run(value)
    set_run(r, 10.5, INK)


def add_table(doc, headers, rows, widths, *, emphasis_last=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, label in zip(header.cells, headers):
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(label)
        set_run(r, 9.5, DARK_BLUE, True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index and len(str(value)) < 22 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run(r, 9.5, INK, emphasis_last and row_index == len(rows) - 1)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
    return table


def add_callout(doc, lead: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(lead)
    set_run(r, 11, DARK_BLUE, True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(body)
    set_run(r, 10.5, INK)


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for level, size, color, before, after in [(1, 16, BLUE, 16, 8), (2, 13, BLUE, 12, 6), (3, 12, DARK_BLUE, 8, 4)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("AgentPlat | Reproducible Empirical Study Appendix")
    set_run(r, 8.5, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Local evidence package V28 | Source commit f041284")
    set_run(r, 8.5, MUTED)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SUPPLEMENTARY METHODS AND RESULTS")
    set_run(r, 10, BLUE, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Reproducible Local Empirical Study: Campaign V28")
    set_run(r, 23, "000000", True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Execution closure, evidence provenance, and analytical eligibility")
    set_run(r, 13, MUTED, False, True)

    add_label_value(doc, "Study identifier", "paper-study-v28")
    add_label_value(doc, "Frozen source commit", "f041284000672b2035138769da7589a6ce89e3f3")
    add_label_value(doc, "Execution environment", "Local Apple Silicon workstation; no cloud services or paid model calls")
    add_label_value(doc, "Collection status", "Complete and collected")

    add_callout(
        doc,
        "Reader-facing conclusion.",
        "The preregistered local execution completed all 48 authorized shards and 960 projections with zero recorded failed executions and USD 0 external spend. The resulting analysis is intentionally labeled ineligible for an empirical claim because three prespecified normative evidence conditions were not satisfied. This is an analytical qualification, not an execution failure.",
    )

    add_heading(doc, "1. Study purpose and scope", 1)
    add_paragraph(doc, "This appendix documents the reproducible local execution of a preregistered comparative simulation campaign. It is designed to make the execution record, evidence topology, and interpretive boundaries suitable for inclusion with a research manuscript. The campaign compares an adaptive collective configuration with a centralized planner across four operating strata and a fixed family of scales and seeds.")

    add_heading(doc, "2. Preregistered execution design", 1)
    add_table(doc, ["Component", "Registered value"], [
        ("Authorized shards", "48"),
        ("Experimental cells", "240"),
        ("Projections", "960"),
        ("Aggregation seed", "20260810"),
        ("Execution policy", "Strictly sequential; stop on failure"),
        ("Source commit", "f041284000672b2035138769da7589a6ce89e3f3"),
        ("Authorization digest", "sha256:d76c1269fd5c98c54662c001565ac8a76089ccc373591c145ed3a2fc6b411625"),
    ], [2500, 6860])
    add_paragraph(doc, "The campaign used content-addressed evidence, immutable shard receipts, and a hash-chained operational event log. Mutable supervisor metadata is treated as operational provenance rather than as a scientific outcome.", after=4)

    add_heading(doc, "3. Execution completion and collection", 1)
    add_table(doc, ["Measure", "Observed result", "Interpretation"], [
        ("Completed shards", "48 / 48", "All authorized shards completed"),
        ("Completed projections", "960 / 960", "Full preregistered projection scope"),
        ("Execution outcomes", "960 succeeded; 0 failed", "No failed projection executions"),
        ("Collection status", "Collected", "Results manifest and analysis emitted"),
        ("External spend", "USD 0", "Local-only execution"),
        ("Recorded shard wall time", "10 h 27 m 38 s", "Aggregate recorded shard duration"),
    ], [2450, 2000, 4910])
    add_paragraph(doc, "The supervisor recorded two recoveries during execution. Each interruption was retained in the operational event chain; only immutable shard receipts count toward completion. The final report confirms 48 completed shards, no missing shards, and 960 completed projections.")

    add_heading(doc, "4. Analysis status and interpretive boundary", 1)
    add_callout(
        doc,
        "Important analytical boundary.",
        "Completeness of execution does not by itself permit a research claim. The collected normative analysis returns the decision ineligible. The report must therefore be cited as an executed, reproducible evidence package rather than as a validated empirical claim of the target capability.",
    )
    add_table(doc, ["Analysis field", "Value"], [
        ("Decision", "ineligible"),
        ("Empirical claim permitted", "No"),
        ("Reason code 1", "normative_role_coherence_horizon_invalid"),
        ("Reason code 2", "normative_role_useful_rate_below_threshold"),
        ("Reason code 3", "normative_convergence_evidence_missing"),
    ], [3000, 6360])
    add_paragraph(doc, "The ineligibility decision preserves the separation between operational reliability and substantive validity. It identifies the next methodological work: satisfy or revise the normative role-coherence horizon, useful-role-rate threshold, and convergence evidence requirements before claiming the evaluated property.")

    add_heading(doc, "5. Evidence provenance", 1)
    add_table(doc, ["Artifact", "Digest or status"], [
        ("Collection manifest", "sha256:5cdd29b44f21bfa10b00aa7eed1023be2826a8ed5322e3159d1a91345fac535b"),
        ("Normative analysis", "sha256:c749f7f7efe7732596dd089cd646fcc97f3d1002785d14301c0ea27ab23ade14"),
        ("Receipt root", "sha256:c74d779e532d6b3930497e0b38e44f718e6345af9fd840842248d4d4dab92c6a"),
        ("Operational events", "102 hash-chained events"),
        ("Final execution report", "Created; 48 per-shard receipt records"),
        ("Content-addressed objects", "3,840 stored objects"),
    ], [3000, 6360])

    add_heading(doc, "6. Reproducibility statement for manuscript use", 1)
    add_paragraph(doc, "The study was executed locally from the frozen commit listed above. The evidence package preserves authorization, source identity, shard completion, projection count, collection manifest, analysis digest, per-shard receipts, and the final supervisor report. Re-execution should use the registered operation and authorization artifacts, retain the fixed aggregation seed, and report analytical eligibility independently from completion status.")

    add_heading(doc, "Suggested manuscript language", 2)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("“We completed a preregistered local campaign comprising 48 shards, 240 experimental cells, and 960 projections from a frozen source commit. The evidence collection completed with no failed projection executions and no external spend. However, the prespecified normative analysis classified the evidence as ineligible for an empirical claim because role-coherence, useful-role-rate, and convergence-evidence criteria were not satisfied. We therefore report the result as an auditable execution record and not as confirmation of the evaluated capability.”")
    set_run(r, 10.5, INK, False, True)

    add_heading(doc, "Appendix A. Artifact inventory", 1)
    add_table(doc, ["Artifact", "Role in the study record"], [
        ("Registered operation", "Frozen study identity and preregistered parameters"),
        ("Authorization record", "Permitted shard scope and collection authorization"),
        ("Content-addressed store", "Immutable projection and evidence content"),
        ("Collection manifest", "Counts, digests, and result-status closure"),
        ("Normative analysis", "Endpoint calculations, decision, and reason codes"),
        ("Final supervisor report", "Execution history, shard receipts, recoveries, and environment summary"),
    ], [3000, 6360])

    doc.core_properties.title = "AgentPlat V28 Reproducible Empirical Study Appendix"
    doc.core_properties.subject = "Supplementary methods and results"
    doc.core_properties.author = "AgentPlat"
    doc.core_properties.comments = "Generated from the local V28 evidence package."
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
